import csv
import io
from datetime import date

from openpyxl import load_workbook

from app.core.permissions import UserRole
from app.models.expense import PaymentMethod
from app.schemas.expense import ExpenseCreate
from app.services import export_service, expense_service
from tests.conftest import make_user


def _make_expense(db_session, tecnico):
    return expense_service.create_expense(
        db_session, actor=tecnico,
        data=ExpenseCreate(
            fecha=date.today(), monto=12345, categoria="repuesto", proveedor="Repuestos Export",
            forma_pago=PaymentMethod.EFECTIVO_PROPIO,
        ),
    )


def test_expenses_to_csv_contiene_encabezado_y_filas(db_session):
    tecnico = make_user(db_session, rol=UserRole.TECNICO, telefono="+56900333001")
    expense = _make_expense(db_session, tecnico)

    content = export_service.expenses_to_csv(db_session, [expense])
    text = content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))

    assert rows[0][0] == "Fecha"
    assert "Repuestos Export" in rows[1]


def test_expenses_to_xlsx_es_un_libro_valido(db_session):
    tecnico = make_user(db_session, rol=UserRole.TECNICO, telefono="+56900333002")
    expense = _make_expense(db_session, tecnico)

    content = export_service.expenses_to_xlsx(db_session, [expense])
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active

    assert ws["A1"].value == "Fecha"
    assert ws.max_row == 2


def test_expenses_to_pdf_produce_un_pdf_valido(db_session):
    tecnico = make_user(db_session, rol=UserRole.TECNICO, telefono="+56900333003")
    expense = _make_expense(db_session, tecnico)

    content = export_service.expenses_to_pdf(db_session, [expense])

    assert content.startswith(b"%PDF")


def test_endpoint_exportar_gastos_csv(client, db_session):
    tecnico = make_user(
        db_session, rol=UserRole.ADMINISTRACION, telefono="+56900333004",
        email="export.test@jacobea.cl", password="clave-segura-123",
    )
    _make_expense(db_session, tecnico)
    db_session.commit()

    login = client.post(
        "/api/auth/login", json={"email": "export.test@jacobea.cl", "password": "clave-segura-123"}
    )
    token = login.json()["access_token"]

    response = client.get(
        "/api/reports/gastos/export", params={"formato": "csv"},
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "Repuestos Export" in response.text


def test_endpoint_exportar_gastos_formato_invalido(client, db_session):
    make_user(
        db_session, rol=UserRole.ADMINISTRACION, telefono="+56900333005",
        email="export2.test@jacobea.cl", password="clave-segura-123",
    )
    db_session.commit()

    login = client.post(
        "/api/auth/login", json={"email": "export2.test@jacobea.cl", "password": "clave-segura-123"}
    )
    token = login.json()["access_token"]

    response = client.get(
        "/api/reports/gastos/export", params={"formato": "docx"},
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 422
